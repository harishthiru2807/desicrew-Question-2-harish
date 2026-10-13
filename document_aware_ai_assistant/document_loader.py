"""PDF, DOCX, and TXT extraction with conservative metadata handling."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

from config import SUPPORTED_EXTENSIONS

HEADING_RE = re.compile(r"^(#{1,6}\s+.+|[A-Z0-9][A-Z0-9 /&(),.'-]{8,}|(\d+(\.\d+)*)\s+[A-Z].+)$")


@dataclass
class PageContent:
    page_number: int | None
    text: str
    headings: list[str] = field(default_factory=list)


@dataclass
class ExtractedDocument:
    filename: str
    file_type: str
    pages: list[PageContent] = field(default_factory=list)
    source_metadata: dict = field(default_factory=dict)
    content_hash: str = ""
    error: str | None = None
    is_empty: bool = False
    needs_ocr: bool = False
    preview_text: str = ""

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text).strip()

    @property
    def ok(self) -> bool:
        return self.error is None and not self.is_empty and not self.needs_ocr and bool(self.full_text)


class DocumentLoadError(Exception):
    """Raised when a file cannot be accepted for indexing."""


def _normalize_filename(filename: str) -> str:
    name = Path(filename).name.strip()
    if not name:
        raise DocumentLoadError("A valid filename is required.")
    return name


def detect_file_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentLoadError(
            f"Unsupported file type '{suffix or 'unknown'}'. Supported types: PDF, DOCX, and TXT."
        )
    return suffix.lstrip(".")


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _looks_like_heading(line: str, font_size: float | None = None, is_bold: bool = False) -> bool:
    text = line.strip()
    if not text or len(text) > 120:
        return False
    if text.endswith("."):
        return False
    if font_size is not None and font_size >= 13.5:
        return True
    if is_bold and len(text) <= 80:
        return True
    return bool(HEADING_RE.match(text)) and len(text.split()) <= 14


def _preview(text: str, limit: int = 1200) -> str:
    cleaned = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[:limit].rstrip() + "…"


def _extract_txt(data: bytes, filename: str) -> ExtractedDocument:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return ExtractedDocument(
            filename=filename,
            file_type="txt",
            error="The TXT file could not be decoded as readable text.",
            is_empty=True,
        )
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return ExtractedDocument(
            filename=filename,
            file_type="txt",
            is_empty=True,
            error="The TXT file does not contain any extractable text.",
        )
    headings = [ln.strip().lstrip("#").strip() for ln in text.split("\n") if _looks_like_heading(ln)]
    page = PageContent(page_number=None, text=text, headings=headings)
    return ExtractedDocument(
        filename=filename,
        file_type="txt",
        pages=[page],
        source_metadata={"encoding_detected": True, "filename": filename},
        content_hash=_sha256(data),
        preview_text=_preview(text),
    )


def _extract_docx(data: bytes, filename: str) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentLoadError("python-docx is required to read DOCX files.") from exc

    try:
        document = Document(io.BytesIO(data))
    except Exception:
        return ExtractedDocument(
            filename=filename,
            file_type="docx",
            error="The DOCX file is invalid or could not be opened.",
            is_empty=True,
        )

    core = document.core_properties
    source_metadata = {
        "filename": filename,
        "author": core.author or None,
        "title": core.title or None,
        "subject": core.subject or None,
        "created": str(core.created) if core.created else None,
    }
    blocks: list[str] = []
    headings: list[str] = []
    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style_name = (para.style.name if para.style is not None else "") or ""
        is_heading = style_name.lower().startswith("heading") or _looks_like_heading(text)
        if is_heading:
            headings.append(text)
            blocks.append(text)
        else:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    full = "\n".join(blocks).strip()
    if not full:
        return ExtractedDocument(
            filename=filename,
            file_type="docx",
            source_metadata=source_metadata,
            is_empty=True,
            error="The DOCX file does not contain any extractable text.",
        )
    page = PageContent(page_number=None, text=full, headings=headings)
    return ExtractedDocument(
        filename=filename,
        file_type="docx",
        pages=[page],
        source_metadata=source_metadata,
        content_hash=_sha256(data),
        preview_text=_preview(full),
    )


def _extract_pdf(data: bytes, filename: str) -> ExtractedDocument:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentLoadError("PyMuPDF is required to read PDF files.") from exc

    try:
        pdf = fitz.open(stream=data, filetype="pdf")
    except Exception:
        return ExtractedDocument(
            filename=filename,
            file_type="pdf",
            error="The PDF file is invalid or could not be opened.",
            is_empty=True,
        )

    meta = pdf.metadata or {}
    source_metadata = {
        "filename": filename,
        "title": meta.get("title") or None,
        "author": meta.get("author") or None,
        "subject": meta.get("subject") or None,
        "page_count": pdf.page_count,
    }
    pages: list[PageContent] = []
    total_chars = 0
    for index in range(pdf.page_count):
        page = pdf[index]
        page_number = index + 1
        headings: list[str] = []
        lines: list[str] = []
        try:
            raw = page.get_text("dict")
        except Exception:
            raw = {"blocks": []}
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                line_text = "".join(span.get("text", "") for span in spans).strip()
                if not line_text:
                    continue
                sizes = [float(span.get("size", 0)) for span in spans]
                fonts = " ".join(span.get("font", "") for span in spans).lower()
                max_size = max(sizes) if sizes else 0.0
                is_bold = "bold" in fonts
                if _looks_like_heading(line_text, font_size=max_size, is_bold=is_bold):
                    headings.append(line_text)
                lines.append(line_text)
        text = "\n".join(lines).strip()
        if not text:
            text = (page.get_text("text") or "").strip()
        total_chars += len(text)
        pages.append(PageContent(page_number=page_number, text=text, headings=headings))
    pdf.close()

    if total_chars == 0:
        return ExtractedDocument(
            filename=filename,
            file_type="pdf",
            pages=pages,
            source_metadata=source_metadata,
            content_hash=_sha256(data),
            is_empty=True,
            needs_ocr=True,
            error=(
                "This PDF contains no extractable text. It may be a scanned image-only document "
                "and would require OCR before indexing. OCR was not performed."
            ),
        )

    preview = _preview("\n\n".join(p.text for p in pages if p.text))
    return ExtractedDocument(
        filename=filename,
        file_type="pdf",
        pages=pages,
        source_metadata=source_metadata,
        content_hash=_sha256(data),
        preview_text=preview,
    )


def load_document(
    file_bytes: bytes,
    filename: str,
    max_bytes: int | None = None,
) -> ExtractedDocument:
    """Validate and extract a single uploaded file."""
    filename = _normalize_filename(filename)
    if max_bytes is not None and len(file_bytes) > max_bytes:
        raise DocumentLoadError(
            f"'{filename}' exceeds the maximum upload size of {max_bytes // (1024 * 1024)} MB."
        )
    if not file_bytes:
        return ExtractedDocument(
            filename=filename,
            file_type=detect_file_type(filename),
            is_empty=True,
            error=f"'{filename}' is empty and cannot be indexed.",
        )

    file_type = detect_file_type(filename)
    if file_type == "txt":
        extracted = _extract_txt(file_bytes, filename)
    elif file_type == "docx":
        extracted = _extract_docx(file_bytes, filename)
    else:
        extracted = _extract_pdf(file_bytes, filename)

    if extracted.ok and not extracted.content_hash:
        extracted.content_hash = _sha256(file_bytes)
    return extracted


def load_from_path(path: str | Path, max_bytes: int | None = None) -> ExtractedDocument:
    path = Path(path)
    data = path.read_bytes()
    return load_document(data, path.name, max_bytes=max_bytes)


def load_from_upload(uploaded_file: BinaryIO, filename: str | None = None, max_bytes: int | None = None) -> ExtractedDocument:
    name = filename or getattr(uploaded_file, "name", "upload.bin")
    data = uploaded_file.read()
    return load_document(data, name, max_bytes=max_bytes)
